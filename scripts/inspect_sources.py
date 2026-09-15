"""Inspect source workbooks without exporting person-level HR data.

The script prints workbook structure and aggregate quality statistics only. It
never modifies a source file and never writes extracted rows to disk.
"""

from __future__ import annotations

import argparse
import re
from collections import Counter, defaultdict
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from pypdf import PdfReader


PERFORMANCE_HEADERS = {
    "Id",
    "Last name",
    "First name",
    "Department",
    "Day",
    "Task",
    "Debut tache",
    "Fin tache",
    "Remark",
    "Customer",
    "Kms",
}
REFERENCE_HEADERS = {
    "Externe referentie",
    "Naam",
    "Woon-werkvervoer uitbetalen",
    "Locatie",
    "Afstand",
    "Bedrag/eenheid",
    "Frequentie",
    "Vervoerswijze",
    "Voertuig wordt effectief gebruikt",
    "Soort abonnement",
    "Tarief periode",
}
ACERTA_HEADERS = {
    "Update code",
    "Extern referentienummer (type HRM-nummer) (*)",
    "Naam werknemer",
    "Loonperiode (*)",
    "Manipulatiecode",
    "Looncode (*)",
    "Eenheden",
    "Bedrag per eenheid",
    "Percentage",
    "Bedrag",
    "Dagen",
    "Kostenplaats",
    "Reden",
    "Startdatum (fractie)",
    "Einddatum (fractie)",
}
NOT_PAY_HEADERS = {
    "Externe referentie",
    "Naam",
    "Volgnummer",
    "Woon-werkvervoer uitbetalen",
    "Afstand",
    "Vervoerswijze",
    "Voertuig wordt effectief gebruikt",
    "Soort abonnement",
    "Tarief periode",
}

SENSITIVE_HEADERS = {
    "id",
    "last name",
    "first name",
    "naam",
    "naam werknemer",
    "externe referentie",
    "extern referentienummer (type hrm-nummer) (*)",
}

SAFE_CATEGORICAL_VALUES = {
    "Remark": {
        "ICTS_OV",
        "ICTS_OV_EXTRA_SHIFT",
        "OV_AC",
        "48h_ICTS_Extra_Shift",
        "OV_DL",
        "Administration",
        "Telework",
        "Agent_OV",
        "Wissel",
        "TRA_NO_SOC_ABO",
        "Verplaatsing",
        "OV_PostNl",
    },
    "Woon-werkvervoer uitbetalen": {
        "Wel uitbetalen",
        "Niet uitbetalen (informatief)",
        "Niet uitbetalen",
        "niet",
        "NIET",
    },
    "Frequentie": {
        "Werkdagen (enkele afstand)",
        "Fiets enkel",
        "Dagelijks fiets (heen en weer)",
        "suppl",
        "alternatief adres",
        "suppl vroeg",
        "auto",
        "vroeg",
        "per dag",
    },
    "Vervoerswijze": {"Privé auto", "Fiets", "Trein", "Dienstwagen", "auto", "Auto"},
    "Voertuig wordt effectief gebruikt": {"Ja"},
    "Soort abonnement": {"Maandabonnement - algemeen"},
    "Tarief periode": {"Weekbedrag/5", "Weekbedrag/6", "Maximale fietsvergoeding"},
}


def clean_header(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip())


def normalized_text(value: Any) -> str:
    if value is None or str(value).strip() == "":
        return ""
    return re.sub(r"\s+", " ", str(value).strip()).casefold()


def cell_type(value: Any) -> str:
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, datetime):
        return "datetime"
    if isinstance(value, date):
        return "date"
    if isinstance(value, time):
        return "time"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    return "text"


def rows_as_values(sheet: Worksheet) -> Iterable[tuple[Any, ...]]:
    return sheet.iter_rows(values_only=True)


def detect_header_row(sheet: Worksheet, expected: set[str]) -> tuple[int | None, list[str]]:
    best_row: int | None = None
    best_headers: list[str] = []
    best_score = 0
    expected_folded = {item.casefold() for item in expected}

    for row_number, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 75), values_only=True),
        start=1,
    ):
        headers = [clean_header(value) for value in row]
        score = len({value.casefold() for value in headers if value} & expected_folded)
        if score > best_score:
            best_row, best_headers, best_score = row_number, headers, score

    minimum = max(2, min(4, len(expected) // 3))
    return (best_row, best_headers) if best_score >= minimum else (None, [])


def header_map(headers: list[str]) -> dict[str, int]:
    return {header: index for index, header in enumerate(headers) if header}


def data_rows(sheet: Worksheet, header_row: int, width: int) -> list[tuple[Any, ...]]:
    rows = []
    for row in sheet.iter_rows(
        min_row=header_row + 1,
        max_row=sheet.max_row,
        max_col=width,
        values_only=True,
    ):
        if any(value is not None and str(value).strip() != "" for value in row):
            rows.append(row)
    return rows


def aggregate_column_quality(
    rows: list[tuple[Any, ...]], headers: list[str]
) -> list[tuple[str, int, str, int]]:
    result = []
    for index, header in enumerate(headers):
        if not header:
            continue
        values = [row[index] if index < len(row) else None for row in rows]
        blanks = sum(value is None or str(value).strip() == "" for value in values)
        types = Counter(
            cell_type(value)
            for value in values
            if value is not None and str(value).strip() != ""
        )
        whitespace = sum(
            isinstance(value, str) and value != value.strip() for value in values
        )
        type_summary = ", ".join(f"{key}:{count}" for key, count in sorted(types.items()))
        result.append((header, blanks, type_summary or "-", whitespace))
    return result


def duplicate_stats(
    rows: list[tuple[Any, ...]], mapping: dict[str, int], keys: list[str]
) -> tuple[int, int]:
    if not all(key in mapping for key in keys):
        return 0, 0
    counts: Counter[tuple[str, ...]] = Counter()
    for row in rows:
        values = tuple(normalized_text(row[mapping[key]]) for key in keys)
        if all(values):
            counts[values] += 1
    duplicate_groups = sum(count > 1 for count in counts.values())
    extra_rows = sum(count - 1 for count in counts.values() if count > 1)
    return duplicate_groups, extra_rows


def normalization_collision_count(values: Iterable[Any]) -> int:
    groups: dict[str, set[str]] = defaultdict(set)
    for value in values:
        if value is None or str(value).strip() == "":
            continue
        raw = str(value)
        groups[normalized_text(raw)].add(raw)
    return sum(len(variants) > 1 for variants in groups.values())


def print_safe_counts(field: str, values: Iterable[Any]) -> None:
    allowed = SAFE_CATEGORICAL_VALUES[field]
    counts: Counter[str] = Counter()
    for value in values:
        cleaned = clean_header(value)
        if not cleaned:
            continue
        if isinstance(value, str) and value.startswith("="):
            counts["<formule>"] += 1
        elif cleaned in allowed:
            counts[cleaned] += 1
        else:
            counts["<overig/onverwacht; inhoud onderdrukt>"] += 1
    rendered = ", ".join(f"{value!r}:{count}" for value, count in counts.most_common())
    print(f"Waarden {field}: {rendered or 'geen'}")


def print_sheet_overview(path: Path, expected: set[str], target_sheet: str) -> None:
    workbook = load_workbook(path, read_only=True, data_only=False)
    print(f"\n## {path.name}")
    print(f"Werkbladen ({len(workbook.sheetnames)}): {', '.join(workbook.sheetnames)}")
    for name in workbook.sheetnames:
        sheet = workbook[name]
        print(f"- {name}: state={sheet.sheet_state}, max_row={sheet.max_row}, max_col={sheet.max_column}")

    if target_sheet not in workbook.sheetnames:
        print(f"WAARSCHUWING: verwacht werkblad ontbreekt: {target_sheet}")
        workbook.close()
        return

    sheet = workbook[target_sheet]
    header_row, headers = detect_header_row(sheet, expected)
    if header_row is None:
        print("WAARSCHUWING: geen betrouwbare header gevonden")
        workbook.close()
        return

    headers = headers[: sheet.max_column]
    mapping = header_map(headers)
    rows = data_rows(sheet, header_row, len(headers))
    print(f"Header: rij {header_row}")
    print(f"Niet-lege datarijen onder header: {len(rows)}")
    print("Kolommen: " + " | ".join(header or "<leeg>" for header in headers))
    missing = sorted(expected - set(headers))
    unexpected = [header for header in headers if header and header not in expected]
    print(f"Verwachte kolommen niet gevonden: {missing or 'geen'}")
    print(f"Extra kolommen: {unexpected or 'geen'}")

    print("Kolomkwaliteit (leeg; types; tekst met randspaties):")
    for header, blanks, types, whitespace in aggregate_column_quality(rows, headers):
        label = "<gevoelig veld>" if header.casefold() in SENSITIVE_HEADERS else header
        print(f"- {label}: leeg={blanks}; types={types}; randspaties={whitespace}")

    if target_sheet == "Total kms":
        id_index = mapping.get("Id")
        summary_rows = 0
        if id_index is not None:
            summary_rows = sum(
                bool(re.search(r"\btotal\b", str(row[id_index]), re.IGNORECASE))
                for row in rows
                if row[id_index] is not None
            )
        duplicate_groups, extra_rows = duplicate_stats(
            rows, mapping, ["Id", "Day", "Customer"]
        )
        print(f"Kandidaat-samenvattingsregels met 'Total' in Id: {summary_rows}")
        print(
            "Kandidaat-duplicaten [Id + Day + Customer]: "
            f"groepen={duplicate_groups}; extra rijen={extra_rows}"
        )
        if "Remark" in mapping:
            print_safe_counts(
                "Remark", (row[mapping["Remark"]] for row in rows)
            )
        if "Customer" in mapping:
            collisions = normalization_collision_count(
                row[mapping["Customer"]] for row in rows
            )
            print(f"Locatienamen met casing/spatievarianten: {collisions}")

    if target_sheet == "Report":
        duplicate_groups, extra_rows = duplicate_stats(
            rows, mapping, ["Externe referentie", "Locatie"]
        )
        print(
            "Herhaalde sleutel [Externe referentie + Locatie]: "
            f"groepen={duplicate_groups}; extra rijen={extra_rows}"
        )
        candidate_records = sum(
            clean_header(row[mapping["Externe referentie"]]) != ""
            and clean_header(row[mapping["Locatie"]]) != ""
            for row in rows
        )
        formula_rows = sum(
            any(isinstance(value, str) and value.startswith("=") for value in row)
            for row in rows
        )
        print(f"Kandidaat-referentieregels met werknemer én locatie: {candidate_records}")
        print(f"Rijen met minstens één formule: {formula_rows}")
        for field in (
            "Woon-werkvervoer uitbetalen",
            "Frequentie",
            "Vervoerswijze",
            "Voertuig wordt effectief gebruikt",
            "Soort abonnement",
            "Tarief periode",
        ):
            if field in mapping:
                print_safe_counts(field, (row[mapping[field]] for row in rows))
        if "Locatie" in mapping:
            collisions = normalization_collision_count(
                row[mapping["Locatie"]] for row in rows
            )
            print(f"Locatienamen met casing/spatievarianten: {collisions}")
            location_rows = [
                row for row in rows if clean_header(row[mapping["Locatie"]])
            ]
            numeric_distances = sum(
                isinstance(row[mapping["Afstand"]], (int, float))
                for row in location_rows
            )
            numeric_amounts = sum(
                isinstance(row[mapping["Bedrag/eenheid"]], (int, float))
                for row in location_rows
            )
            blank_references = sum(
                not clean_header(row[mapping["Externe referentie"]])
                for row in location_rows
            )
            blank_names = sum(
                not clean_header(row[mapping["Naam"]]) for row in location_rows
            )
            keyword_counts = {
                keyword: sum(
                    keyword in normalized_text(row[mapping["Locatie"]])
                    for row in location_rows
                )
                for keyword in ("vroeg", "laat", "48", "suppl")
            }
            locations_by_name: dict[str, set[str]] = defaultdict(set)
            for row in location_rows:
                name = normalized_text(row[mapping["Naam"]])
                location = normalized_text(row[mapping["Locatie"]])
                if name and location:
                    locations_by_name[name].add(location)
            base_and_early_pairs = 0
            names_with_early = 0
            for locations in locations_by_name.values():
                if any("vroeg" in location for location in locations):
                    names_with_early += 1
                for location in locations:
                    if "vroeg" not in location:
                        continue
                    base = re.sub(r"\s+vroeg(?:\s+.*)?$", "", location).strip()
                    if base in locations:
                        base_and_early_pairs += 1
            print(
                "Locatierijen: "
                f"totaal={len(location_rows)}; numerieke afstand={numeric_distances}; "
                f"numeriek bedrag={numeric_amounts}; zonder externe referentie={blank_references}; "
                f"zonder naam={blank_names}"
            )
            print(
                "Locatie-keywords: "
                + "; ".join(f"{key}={value}" for key, value in keyword_counts.items())
            )
            print(
                "Basis plus vroeg-combinaties: "
                f"paren={base_and_early_pairs}; naamgroepen met vroeg={names_with_early}"
            )

    workbook.close()


def print_acerta_overview(path: Path) -> None:
    workbook = load_workbook(path, read_only=True, data_only=False)
    print(f"\n## {path.name}")
    print(f"Werkbladen ({len(workbook.sheetnames)}): {', '.join(workbook.sheetnames)}")
    for name in workbook.sheetnames:
        sheet = workbook[name]
        expected = ACERTA_HEADERS if name.strip() == "Afwijkende loonelementen" else NOT_PAY_HEADERS
        header_row, headers = detect_header_row(sheet, expected)
        print(f"- {name}: state={sheet.sheet_state}, max_row={sheet.max_row}, max_col={sheet.max_column}")
        if header_row is not None:
            rows = data_rows(sheet, header_row, sheet.max_column)
            print(f"  Header: rij {header_row}; niet-lege datarijen: {len(rows)}")
            print("  Kolommen: " + " | ".join(headers[: sheet.max_column]))
            print("  Kolomkwaliteit (leeg; types; tekst met randspaties):")
            for header, blanks, types, whitespace in aggregate_column_quality(
                rows, headers[: sheet.max_column]
            ):
                label = "<gevoelig veld>" if header.casefold() in SENSITIVE_HEADERS else header
                print(f"  - {label}: leeg={blanks}; types={types}; randspaties={whitespace}")
        else:
            labels_only = []
            for row in sheet.iter_rows(
                min_row=1,
                max_row=min(sheet.max_row, 10),
                max_col=sheet.max_column,
                values_only=True,
            ):
                labels = [clean_header(value) for value in row if clean_header(value)]
                if labels:
                    labels_only.append(labels)
            print(
                "  Geen betrouwbare header gedetecteerd; "
                f"niet-lege rijen in eerste 10 rijen={len(labels_only)} "
                "(inhoud onderdrukt voor privacy)"
            )
    workbook.close()


def print_pdf_overview(path: Path) -> None:
    reader = PdfReader(path)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    terms = [
        "120%",
        "150%",
        "0,31",
        "0,25",
        "0,4326",
        "0,36",
        "fietsvergoeding",
        "atypische uurroosters",
    ]
    print(f"\n## {path.name}")
    print(f"Pagina's: {len(reader.pages)}")
    print(f"Geëxtraheerde tekens: {len(text)}")
    print("Gevonden relevante termen: " + ", ".join(term for term in terms if term.casefold() in text.casefold()))


def print_cross_source_overview(performance_path: Path, reference_path: Path) -> None:
    performance_workbook = load_workbook(
        performance_path, read_only=True, data_only=True
    )
    performance_sheet = performance_workbook["Total kms"]
    performance_header_row, performance_headers = detect_header_row(
        performance_sheet, PERFORMANCE_HEADERS
    )
    assert performance_header_row is not None
    performance_mapping = header_map(performance_headers)
    performance_rows = data_rows(
        performance_sheet, performance_header_row, len(performance_headers)
    )

    performance_ids: set[str] = set()
    first_last: dict[str, set[str]] = defaultdict(set)
    last_first: dict[str, set[str]] = defaultdict(set)
    for row in performance_rows:
        employee_id = normalized_text(row[performance_mapping["Id"]])
        if not employee_id or "total" in employee_id:
            continue
        first_name = normalized_text(row[performance_mapping["First name"]])
        last_name = normalized_text(row[performance_mapping["Last name"]])
        performance_ids.add(employee_id)
        first_last[normalized_text(f"{first_name} {last_name}")].add(employee_id)
        last_first[normalized_text(f"{last_name} {first_name}")].add(employee_id)
    performance_workbook.close()

    reference_workbook = load_workbook(reference_path, read_only=True, data_only=True)
    reference_sheet = reference_workbook["Report"]
    reference_header_row, reference_headers = detect_header_row(
        reference_sheet, REFERENCE_HEADERS
    )
    assert reference_header_row is not None
    reference_mapping = header_map(reference_headers)
    reference_rows = data_rows(
        reference_sheet, reference_header_row, len(reference_headers)
    )

    external_references: set[str] = set()
    reference_names: dict[str, set[str]] = defaultdict(set)
    for row in reference_rows:
        external_reference = normalized_text(
            row[reference_mapping["Externe referentie"]]
        )
        name = normalized_text(row[reference_mapping["Naam"]])
        if external_reference:
            external_references.add(external_reference)
        if external_reference and name:
            reference_names[name].add(external_reference)
    reference_workbook.close()

    first_last_overlap = set(first_last) & set(reference_names)
    last_first_overlap = set(last_first) & set(reference_names)
    deterministic_last_first = sum(
        len(last_first[name]) == 1 and len(reference_names[name]) == 1
        for name in last_first_overlap
    )

    print("\n## Kruisbronmatching (alleen geaggregeerd)")
    print(f"Unieke prestatie-ID's: {len(performance_ids)}")
    print(f"Unieke externe referenties: {len(external_references)}")
    print(f"Exacte ID-overlap: {len(performance_ids & external_references)}")
    print(f"Unieke prestatienamen: {len(first_last)}")
    print(f"Unieke referentienamen op ID-rijen: {len(reference_names)}")
    print(f"Naamoverlap First + Last: {len(first_last_overlap)}")
    print(f"Naamoverlap Last + First: {len(last_first_overlap)}")
    print(f"Deterministische Last + First-overlap: {deterministic_last_first}")
    print(
        "Referentienamen met meerdere externe referenties: "
        f"{sum(len(values) > 1 for values in reference_names.values())}"
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--performance", type=Path, required=True)
    parser.add_argument("--reference", type=Path, required=True)
    parser.add_argument("--acerta", type=Path, required=True)
    parser.add_argument("--pdf", type=Path)
    return parser.parse_args()


def require_files(paths: Iterable[Path | None]) -> None:
    missing = [str(path) for path in paths if path is not None and not path.is_file()]
    if missing:
        raise SystemExit("Bronbestand(en) niet gevonden: " + ", ".join(missing))


def main() -> None:
    args = parse_args()
    require_files((args.performance, args.reference, args.acerta, args.pdf))
    print_sheet_overview(args.performance, PERFORMANCE_HEADERS, "Total kms")
    print_sheet_overview(args.reference, REFERENCE_HEADERS, "Report")
    print_acerta_overview(args.acerta)
    if args.pdf:
        print_pdf_overview(args.pdf)
    print_cross_source_overview(args.performance, args.reference)


if __name__ == "__main__":
    main()
