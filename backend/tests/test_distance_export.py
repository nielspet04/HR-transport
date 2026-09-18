from io import BytesIO
import json
import pytest
from openpyxl import load_workbook
from app import routing
from app.distance_export import workbook_bytes,export_rows,shift_label
from app.distances import whole_kms
from test_routing import prepared


@pytest.mark.parametrize('raw,expected',[('17.00001',18),('17.00',17),('0',0),('0.0001',1),('17,2',18),('60.01',61)])
def test_ceiling(raw,expected):assert whole_kms(raw)==expected


@pytest.mark.parametrize('raw',[None,True,'NaN','Infinity','-1',''])
def test_invalid_distance_not_zero(raw):
    with pytest.raises(ValueError):whole_kms(raw)


def test_night_shift_without_explicit_offset():
    assert shift_label({'start':'19:00','end':'07:00','end_day_offset':None})=='19:00 – 07:00 (+1 dag)'


def export_fixture(tmp_path,monkeypatch):
    store,run,p,wid,_=prepared(tmp_path)
    m={**p['movements'][0],'id':1,'source_shifts':[{'row':2,'customer':'TUI','task':'Controle','start':'22:00','end':'06:00','end_day_offset':1}]}
    with store.transaction() as db:
        p={**p,'movements':[m]}
        db.execute('UPDATE matching_runs SET payload=? WHERE id=?',(json.dumps(p),run))
        store.insert_matching(db,{**p,'month':'2026-09','movements':[{**m,'day':'2026-09-01'}]},'/fictional.xlsx')
    monkeypatch.setattr(routing,'token_value',lambda:'pk.fake')
    monkeypatch.setattr(routing,'request_distance',lambda *a:{'meters':'17000.1','kms':'17.0001','alternatives_count':1})
    route=store.get_route_plan(run)['routes'][0]
    store.apply('route_distance_request',{'run_id':run,'cache_key':route['cache_key'],'consent':True},store.snapshot()['revision'])
    return store,run


def test_all_months_and_modes_literal_names_typed_whole_distances(tmp_path,monkeypatch):
    store,run=export_fixture(tmp_path,monkeypatch)
    revision=store.snapshot()['revision']
    monkeypatch.setattr(routing,'request_distance',lambda *a:pytest.fail('Export must not request Mapbox'))
    result=workbook_bytes(store,run)
    wb=load_workbook(BytesIO(result),data_only=False)
    assert wb.sheetnames==['Afstanden','Shiften']
    assert wb['Afstanden']['C6'].value==18 and wb['Afstanden']['C6'].data_type=='n'
    ws=wb['Shiften'];rows=list(ws.iter_rows(min_row=6,values_only=True))
    assert len(rows)==6 and {r[4].month for r in rows}=={8,9}
    assert all(r[3]=='22:00 – 06:00 (+1 dag)' for r in rows)
    assert all(r[2]==18 for r in rows if r[1]=='Auto')
    assert all(r[2] is None and r[9]=='AFSTAND ONTBREEKT' for r in rows if r[1]=='Fiets')
    assert all(r[2]=='n.v.t.' and r[9]=='N.V.T.' for r in rows if r[1]=='Trein')
    assert all(r[11]=='Vervoerskeuze controleren' for r in rows)
    assert ws.freeze_panes=='C6' and len(ws.tables)==1
    assert store.snapshot()['revision']==revision
    with store.connect() as db:assert db.execute('SELECT kms FROM route_distances').fetchone()[0]=='17.0001'


def test_historical_distance_not_omitted(tmp_path,monkeypatch):
    store,run=export_fixture(tmp_path,monkeypatch)
    with store.transaction() as db:
        r=dict(db.execute('SELECT * FROM route_distances').fetchone());r.update(cache_key='historical',id=999,kms='20.1')
        db.execute('INSERT INTO route_distances('+','.join(r)+') VALUES('+','.join('?' for _ in r)+')',list(r.values()))
    distances,shifts,months=export_rows(store,run)
    assert {r[2] for r in distances}=={18,21}


def test_source_names_and_formula_injection(tmp_path,monkeypatch):
    store,run=export_fixture(tmp_path,monkeypatch)
    with store.transaction() as db:db.execute('UPDATE workers SET name=?',('=HYPERLINK("bad")',))
    # Refresh digest so that the intentionally changed config is current.
    from app.matching import configuration_digest
    with store.transaction() as db:
        digest=configuration_digest(store.matching_config(db))
        for row in list(db.execute('SELECT id,payload FROM matching_runs')):
            p=json.loads(row['payload']);p['configuration_digest']=digest
            db.execute('UPDATE matching_runs SET payload=? WHERE id=?',(json.dumps(p),row['id']))
    wb=load_workbook(BytesIO(workbook_bytes(store,run)))
    assert wb['Afstanden']['A6'].data_type=='s'
    assert wb['Afstanden']['A6'].value.startswith("'=")
