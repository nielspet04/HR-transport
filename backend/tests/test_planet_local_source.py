"""Opt-in check of the new snapshot; never stores/displays person-level rows."""

import os
from collections import Counter
from hashlib import sha256
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.importers.planet import import_planet


@pytest.mark.skipif(not os.environ.get("PLANET_SOURCE"), reason="Set PLANET_SOURCE for the private local snapshot check")
def test_clean_planet_snapshot_row_counts_and_month_selection():
    path = Path(os.environ["PLANET_SOURCE"])
    before = sha256(path.read_bytes()).hexdigest()
    result = import_planet(path)
    assert result.report.imported_rows == result.report.data_rows_seen == 9472
    assert result.report.rejected_rows == result.report.summary_rows == 0
    assert not result.has_errors
    assert result.report.source_sha256 == before
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        raw_month_counts = Counter()
        agreements = []
        for shift, row in zip(result.shifts, workbook["Total kms"].iter_rows(min_row=2, values_only=True)):
            raw_month_counts[row[4][:7]] += 1
            agreements.append(
                shift.employee_id == row[0]
                and shift.day.isoformat() == row[4]
                and shift.start_time.isoformat(timespec="minutes") == row[6]
                and ("24:00" if shift.end_time_day_offset == 1 else shift.end_time.isoformat(timespec="minutes")) == row[7]
                and shift.customer == row[9]
                and shift.remark == (row[8].strip() if row[8] else None)
            )
        assert all(agreements), "Core source fields must agree for every imported row"
    finally:
        workbook.close()
    assert Counter(shift.day.strftime("%Y-%m") for shift in result.shifts) == raw_month_counts
    august = import_planet(path, month="2026-08")
    assert august.report.imported_rows == raw_month_counts["2026-08"]
    assert august.report.outside_month_rows == 9472 - raw_month_counts["2026-08"]
    assert sha256(path.read_bytes()).hexdigest() == before
