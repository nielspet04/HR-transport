from io import BytesIO
import json

from openpyxl import load_workbook

from app.configuration.addresses import append_address
from app.configuration.routes import RouteStore
from app.employee_distance_export import HEADERS, export_rows, workbook_bytes
from app.geocoding import fingerprint
from app.routing import POLICY


def prepared(tmp_path):
    store = RouteStore(tmp_path/'routes.sqlite3')
    store.apply('route_add', {'name':'Voorbeeld Alex','location':'LUCHTHAVEN','mode':'Privé auto',
        'kms':'12.5','valid_from':'2026-01-01','reason':'Test'}, store.snapshot()['revision'])
    worker = store.snapshot()['workers'][0]
    home = {'street':'Teststraat','number':'12','unit':'A','postal_code':'1000','city':'Brussel','country':'BE'}
    site = {**home, 'street':'Luchthavenlaan', 'number':'1', 'unit':''}
    with store.transaction() as db:
        append_address(db, worker['id'], '2026-01-01', home, 'Test')
        db.execute('''INSERT INTO location_address_versions
            (location_key,location,valid_from,address,changed_at,reason) VALUES(?,?,?,?,?,?)''',
            ('luchthaven','LUCHTHAVEN','2026-01-01',json.dumps(site),'test','Test'))
        db.execute('''INSERT INTO route_distances
            (cache_key,profile,origin_hash,destination_hash,origin_geocode_id,destination_geocode_id,
             policy,status,meters,kms,requested_at,permission_basis)
             VALUES(?,?,?,?,?,?,?,?,?,?,?,?)''',
            ('route','driving',fingerprint(home),fingerprint(site),1,2,POLICY,'READY','17200','17.2','test','Test'))
    return store, worker


def test_export_compares_current_mapbox_and_manual_distance(tmp_path):
    store, worker = prepared(tmp_path)
    rows = export_rows(store, '2026-09-29')
    assert rows == [['Voorbeeld Alex','Teststraat 12 bus A, 1000 Brussel, BE',18,12.5,'LUCHTHAVEN','Privé auto']]
    workbook = load_workbook(BytesIO(workbook_bytes(store)))
    sheet = workbook['Werknemerafstanden']
    assert tuple(cell.value for cell in sheet[3]) == HEADERS
    assert sheet['A4'].value == 'Voorbeeld Alex'
    assert sheet['C4'].value == 18 and sheet['D4'].value == 12.5
    assert len(sheet.tables) == 1 and sheet.freeze_panes == 'C4'
