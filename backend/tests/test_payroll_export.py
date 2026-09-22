from datetime import date
from io import BytesIO
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from app.payroll_export import HEADERS, payroll_rows, replace_example_rows, workbook_bytes


def payload():
    movements = [
        {'id': 1, 'planet_id': 'p1', 'day': '2026-08-01', 'location': 'LUCHTHAVEN'},
        {'id': 2, 'planet_id': 'p1', 'day': '2026-08-02', 'location': 'LUCHTHAVEN'},
        {'id': 3, 'planet_id': 'p1', 'day': '2026-08-03', 'location': 'LUCHTHAVEN'},
        {'id': 4, 'planet_id': 'p1', 'day': '2026-08-04', 'location': 'POSTNL'},
        {'id': 5, 'planet_id': 'p1', 'day': '2026-08-05', 'location': 'POSTNL'},
        {'id': 6, 'planet_id': 'p1', 'day': '2026-08-06', 'location': 'POSTNL'},
        {'id': 7, 'planet_id': 'p1', 'day': '2026-08-07', 'location': 'POSTNL'},
    ]
    rows = [
        {'movement_id': 1, 'status': 'CALCULATED', 'amount': '6.70', 'tariff_kind': 'STANDARD'},
        {'movement_id': 2, 'status': 'CALCULATED', 'amount': '6.70', 'tariff_kind': 'STANDARD'},
        {'movement_id': 3, 'status': 'CALCULATED', 'amount': '7.50', 'tariff_kind': 'SPECIAL'},
        {'movement_id': 4, 'status': 'CALCULATED', 'amount': '7.50', 'tariff_kind': 'EXTRA48'},
        {'movement_id': 5, 'status': 'CALCULATED', 'amount': '7.50', 'tariff_kind': 'SPECIAL'},
        {'movement_id': 6, 'status': 'CALCULATED', 'amount': '8.14', 'tariff_kind': 'BICYCLE'},
        {'movement_id': 7, 'status': 'EXCLUDED_TRAIN', 'amount': None},
    ]
    return {'month': '2026-08', 'stale': False,
        'agents': [{'planet_id': 'p1', 'worker_id': 10}], 'movements': movements,
        'calculation': {'rows': rows, 'monthly': {'ready': True, 'external_references_ready': True,
            'calculated_total': '44.04', 'employees': [{'worker_id': 10, 'name': 'Voorbeeld Alex',
                'external_reference': '00123', 'external_reference_status': 'READY'}]}}}


def template_bytes():
    workbook = Workbook(); sheet = workbook.active; sheet.title = 'Afwijkende loonelementen'
    sheet.append(list(HEADERS)); sheet.append([None, 'c', '00999', 'Bestaand', date(2026, 7, 1), None,
        '25', 1, 3.43, None, None, None, None, 'Bestaand', date(2026, 7, 1), date(2026, 7, 31)])
    sheet['C2'].number_format = '@'; sheet['E2'].number_format = 'd/mm/yyyy;@'; sheet['I2'].number_format = '###0.00000'
    sheet.auto_filter.ref = 'A1:XFD2'
    other = workbook.create_sheet('LIST NIET UITBETALEN '); other['A1'] = 'Bewaren'; other['A2'] = 'Ongewijzigd'
    stream = BytesIO(); workbook.save(stream); workbook.close(); return stream.getvalue()


def test_groups_by_worker_location_code_and_exact_amount():
    rows = payroll_rows(payload())
    assert [(row[6], row[7], row[8], row[13]) for row in rows] == [
        ('25', 2, 6.7, 'LUCHTHAVEN'), ('26', 1, 7.5, 'LUCHTHAVEN'),
        ('26', 2, 7.5, 'POSTNL'), ('420', 1, 8.14, 'POSTNL')]
    assert all(row[2] == '00123' and row[4] == date(2026, 8, 1) for row in rows)
    assert all(row[14] == date(2026, 8, 1) and row[15] == date(2026, 8, 31) for row in rows)


@pytest.mark.parametrize('change,message', [
    (lambda data: data.update(stale=True), 'Vernieuw'),
    (lambda data: data['calculation']['monthly'].update(ready=False), 'berekeningsblokkeringen'),
    (lambda data: data['calculation']['monthly'].update(external_references_ready=False), 'loonnummer'),
    (lambda data: data['calculation']['rows'][0].update(tariff_kind='UNKNOWN'), 'vergoedingssoort'),
])
def test_safety_blocks_incomplete_or_unknown_month(change, message):
    data = payload(); change(data)
    with pytest.raises(ValueError, match=message): payroll_rows(data)


def test_replaces_examples_and_preserves_template_structure_and_styles():
    original = template_bytes(); rows = payroll_rows(payload()); result = replace_example_rows(original, rows)
    before = load_workbook(BytesIO(original)); after = load_workbook(BytesIO(result))
    assert after.sheetnames == before.sheetnames == ['Afwijkende loonelementen', 'LIST NIET UITBETALEN ']
    assert [after['Afwijkende loonelementen'].cell(1, c).value for c in range(1, 17)] == list(HEADERS)
    assert list(after['LIST NIET UITBETALEN '].values) == list(before['LIST NIET UITBETALEN '].values)
    sheet = after['Afwijkende loonelementen']; assert sheet.max_row == 1 + len(rows)
    assert all(row[2].value != '00999' for row in sheet.iter_rows(min_row=2))
    assert sheet['C2'].value == '00123' and sheet['C2'].data_type == 's'
    assert sheet['H2'].value == 2 and sheet['H2'].data_type == 'n'
    assert sheet['I2'].value == 6.7 and sheet['I2'].data_type == 'n'
    assert sheet['E2'].value.date() == date(2026, 8, 1)
    assert sheet.auto_filter.ref == f'A1:XFD{sheet.max_row}'
    with ZipFile(BytesIO(original)) as source, ZipFile(BytesIO(result)) as exported:
        assert source.namelist() == exported.namelist()
        assert all(source.read(name) == exported.read(name) for name in source.namelist()
                   if name != 'xl/worksheets/sheet1.xml')


def test_formula_like_name_and_location_remain_literal_text():
    data = payload(); data['calculation']['monthly']['employees'][0]['name'] = '=HYPERLINK("bad")'
    for movement in data['movements']:
        movement['location'] = '+TEST'
    result = replace_example_rows(template_bytes(), payroll_rows(data))
    sheet = load_workbook(BytesIO(result), data_only=False)['Afwijkende loonelementen']
    assert sheet['D2'].value == '=HYPERLINK("bad")' and sheet['D2'].data_type == 's'
    assert sheet['N2'].value == '+TEST' and sheet['N2'].data_type == 's'


def test_fixed_local_template_is_used_without_monthly_upload(tmp_path):
    class Store:
        path = tmp_path / 'routes.sqlite3'
        def get_matching(self, run_id):
            assert run_id == 7
            return payload()
    (tmp_path / 'acerta-template.xlsx').write_bytes(template_bytes())
    result, month = workbook_bytes(Store(), 7)
    sheet = load_workbook(BytesIO(result))['Afwijkende loonelementen']
    assert month == '2026-08' and sheet.max_row == 5 and sheet['D2'].value == 'Voorbeeld Alex'


def test_missing_fixed_template_is_explicit(tmp_path):
    class Store:
        path = tmp_path / 'routes.sqlite3'
    with pytest.raises(ValueError, match='vaste Accerta-sjabloon'): workbook_bytes(Store(), 7)
