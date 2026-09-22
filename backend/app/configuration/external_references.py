"""Dated external payroll references attached to existing workers."""
from collections import Counter, defaultdict
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile
import re

from openpyxl import load_workbook

from app.importers.reference import norm
from .store import required, valid_day


SCHEMA = '''
CREATE TABLE IF NOT EXISTS employee_external_references(
 id INTEGER PRIMARY KEY, worker_id INTEGER NOT NULL REFERENCES workers(id),
 valid_from TEXT NOT NULL, external_reference TEXT NOT NULL,
 changed_at TEXT NOT NULL, reason TEXT NOT NULL,
 source_row INTEGER, source_hash TEXT
);
CREATE INDEX IF NOT EXISTS employee_external_references_worker_date
 ON employee_external_references(worker_id,valid_from,id);
CREATE TABLE IF NOT EXISTS external_reference_imports(
 hash TEXT PRIMARY KEY, valid_from TEXT NOT NULL, report TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS external_reference_candidates(
 id INTEGER PRIMARY KEY, import_hash TEXT NOT NULL REFERENCES external_reference_imports(hash),
 source_row INTEGER NOT NULL, name TEXT NOT NULL, external_reference TEXT NOT NULL,
 status TEXT NOT NULL, worker_id INTEGER REFERENCES workers(id)
);
'''

HEADERS = ('Naam', 'Externe referentie')


def validate_reference(value):
    value = required(value)
    if not re.fullmatch(r'\d{1,50}', value):
        raise ValueError('Extern loonreferentienummer moet uitsluitend uit cijfers bestaan.')
    return value


def read_external_references(source):
    """Read only the two explicitly required columns and preserve identifiers as text."""
    source = Path(source)
    if source.stat().st_size > 5_000_000:
        raise ValueError('Loonreferentiebestand is groter dan 5 MB.')
    digest = sha256(source.read_bytes()).hexdigest()
    with ZipFile(source) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(item.file_size for item in entries) > 64_000_000:
            raise ValueError('Loonreferentiebestand bevat te veel uitgepakte gegevens.')
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        if 'Report' not in workbook.sheetnames:
            raise ValueError('Werkblad Report ontbreekt.')
        sheet = workbook['Report']; sheet.reset_dimensions()
        rows = list(sheet.iter_rows())
        mapping = None; header_row = None
        for index, row in enumerate(rows[:75]):
            names = [norm(cell.value) for cell in row]
            if all(norm(header) in names for header in HEADERS):
                if any(names.count(norm(header)) != 1 for header in HEADERS):
                    raise ValueError('Dubbele vereiste kolom in loonreferentiebestand.')
                mapping = {header: names.index(norm(header)) for header in HEADERS}
                header_row = index + 1
                break
        if mapping is None:
            raise ValueError('Kolommen Naam en Externe referentie ontbreken.')
        records = []
        for number, row in enumerate(rows[header_row:], header_row + 1):
            if not any(cell.value not in (None, '') for cell in row):
                continue
            name_cell = row[mapping['Naam']]
            reference_cell = row[mapping['Externe referentie']]
            if name_cell.data_type in ('f', 'e') or reference_cell.data_type in ('f', 'e'):
                raise ValueError(f'Bronrij {number}: formules en Excel-fouten zijn niet toegestaan.')
            name = required(name_cell.value)
            if not isinstance(reference_cell.value, str):
                raise ValueError(f'Bronrij {number}: bewaar het loonnummer als tekst zodat voorloopnullen behouden blijven.')
            reference = validate_reference(reference_cell.value)
            records.append({'row': number, 'name': name, 'external_reference': reference})
        if not records:
            raise ValueError('Geen loonreferenties gevonden.')
        names = Counter(norm(record['name']) for record in records)
        references = Counter(record['external_reference'] for record in records)
        if any(count > 1 for count in names.values()):
            raise ValueError('Een werknemer staat meermaals in het loonreferentiebestand.')
        if any(count > 1 for count in references.values()):
            raise ValueError('Een extern loonreferentienummer staat bij meerdere werknemers.')
        return digest, records
    finally:
        workbook.close()


def append_reference(db, worker_id, start, reference, reason, source_row=None, source_hash=None):
    if type(worker_id) is not int or not db.execute('SELECT 1 FROM workers WHERE id=?', (worker_id,)).fetchone():
        raise ValueError('Kies een bestaande werknemer.')
    start = valid_day(start); reference = validate_reference(reference); reason = required(reason)
    last = db.execute('SELECT max(valid_from) FROM employee_external_references WHERE worker_id=?', (worker_id,)).fetchone()[0]
    if last and start < last:
        raise ValueError('Kies de laatste ingangsdatum of een latere datum; oude loonnummers blijven bewaard.')
    return db.execute('''INSERT INTO employee_external_references
        (worker_id,valid_from,external_reference,changed_at,reason,source_row,source_hash)
        VALUES(?,?,?,?,?,?,?)''', (worker_id, start, reference, datetime.now(timezone.utc).isoformat(),
        reason, source_row, source_hash)).lastrowid


def _confirmed_name_aliases(db):
    aliases = defaultdict(set)
    for row in db.execute('SELECT id,name_key FROM workers'):
        aliases[row['name_key']].add(row['id'])
    # These names were already reviewed during the address import. Reusing them is
    # safer than introducing a second independent fuzzy-name decision.
    for row in db.execute("SELECT name,worker_id FROM address_candidates WHERE status='LINKED' AND worker_id IS NOT NULL"):
        aliases[norm(row['name'])].add(row['worker_id'])
    return aliases


def import_external_references(store, source, start, revision):
    start = valid_day(start)
    digest, records = read_external_references(source)
    with store.transaction(revision) as db:
        old = db.execute('SELECT valid_from,report FROM external_reference_imports WHERE hash=?', (digest,)).fetchone()
        if old:
            if old['valid_from'] != start:
                raise ValueError('Dit bestand is al ingelezen met een andere ingangsdatum.')
            import json
            return {**json.loads(old['report']), 'already_imported': True}
        aliases = _confirmed_name_aliases(db)
        report = {'source_rows': len(records), 'linked': 0, 'review': 0}
        db.execute('INSERT INTO external_reference_imports(hash,valid_from,report) VALUES(?,?,?)', (digest, start, '{}'))
        for record in records:
            matches = aliases.get(norm(record['name']), set())
            worker_id = next(iter(matches)) if len(matches) == 1 else None
            existing = worker_id and db.execute('SELECT 1 FROM employee_external_references WHERE worker_id=?', (worker_id,)).fetchone()
            status = 'LINKED' if worker_id and not existing else 'REVIEW'
            if status == 'LINKED':
                append_reference(db, worker_id, start, record['external_reference'],
                    'Excel EXTERN REF NR · ingangsdatum bevestigd', record['row'], digest)
                report['linked'] += 1
            else:
                report['review'] += 1
            db.execute('''INSERT INTO external_reference_candidates
                (import_hash,source_row,name,external_reference,status,worker_id) VALUES(?,?,?,?,?,?)''',
                (digest, record['row'], record['name'], record['external_reference'], status, worker_id))
        import json
        db.execute('UPDATE external_reference_imports SET report=? WHERE hash=?', (json.dumps(report), digest))
    return report


def effective(rows, worker_id, day):
    available = [row for row in rows if row['worker_id'] == worker_id and row['valid_from'] <= day]
    return max(available, key=lambda row: (row['valid_from'], row['id'])) if available else None


def snapshot(db):
    return {
        'external_references': [dict(row) for row in db.execute(
            'SELECT * FROM employee_external_references ORDER BY valid_from,id')],
        'external_reference_candidates': [dict(row) for row in db.execute(
            'SELECT * FROM external_reference_candidates ORDER BY id')],
    }


def apply(db, action, data):
    if action == 'external_reference_save':
        return append_reference(db, data.get('worker_id'), data.get('valid_from'),
            data.get('external_reference'), data.get('reason'))
    candidate = db.execute('SELECT c.*,i.valid_from FROM external_reference_candidates c '
        'JOIN external_reference_imports i ON i.hash=c.import_hash WHERE c.id=?',
        (data.get('candidate_id'),)).fetchone()
    if not candidate or candidate['status'] != 'REVIEW':
        raise ValueError('Deze loonreferentiebronregel is niet meer beschikbaar.')
    reason = required(data.get('reason'))
    if action == 'external_reference_link':
        worker_id = data.get('worker_id')
        append_reference(db, worker_id, candidate['valid_from'], candidate['external_reference'],
            reason, candidate['source_row'], candidate['import_hash'])
        db.execute("UPDATE external_reference_candidates SET status='LINKED',worker_id=? WHERE id=?",
            (worker_id, candidate['id']))
    elif action == 'external_reference_ignore':
        db.execute("UPDATE external_reference_candidates SET status='IGNORED' WHERE id=?", (candidate['id'],))
    else:
        raise ValueError('Onbekende loonreferentieactie.')


def enrich_monthly(db, payload, monthly):
    """Attach the reference valid on each shift date without changing arithmetic readiness."""
    rows = [dict(row) for row in db.execute('SELECT * FROM employee_external_references ORDER BY valid_from,id')]
    movement_days = defaultdict(set)
    agents = {agent['planet_id']: agent for agent in payload['agents']}
    for movement in payload['movements']:
        worker_id = agents[movement['planet_id']].get('worker_id')
        if worker_id:
            movement_days[worker_id].add(movement['day'])
    missing = changed = 0
    for employee in monthly['employees']:
        references = {item['external_reference'] for day in movement_days.get(employee['worker_id'], set())
            if (item := effective(rows, employee['worker_id'], day))}
        uncovered = any(effective(rows, employee['worker_id'], day) is None
            for day in movement_days.get(employee['worker_id'], set()))
        if uncovered or not references:
            employee.update(external_reference=None, external_reference_status='MISSING')
            missing += 1
        elif len(references) > 1:
            employee.update(external_reference=None, external_reference_status='CHANGED_DURING_MONTH',
                external_references=sorted(references))
            changed += 1
        else:
            employee.update(external_reference=next(iter(references)), external_reference_status='READY')
    monthly.update(external_reference_missing=missing, external_reference_changed=changed,
        external_references_ready=missing == 0 and changed == 0)
    return monthly
