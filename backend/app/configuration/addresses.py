"""Dated addresses attached to existing workers; no external API calls."""
from collections import Counter
from datetime import datetime, timezone
from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from app.importers.reference import norm
from .store import required, valid_day

SCHEMA = '''
CREATE TABLE IF NOT EXISTS worker_addresses(
 id INTEGER PRIMARY KEY, worker_id INTEGER NOT NULL REFERENCES workers(id),
 valid_from TEXT NOT NULL, address TEXT NOT NULL, changed_at TEXT NOT NULL,
 reason TEXT NOT NULL, source_row INTEGER, source_hash TEXT);
CREATE TABLE IF NOT EXISTS address_imports(
 hash TEXT PRIMARY KEY, valid_from TEXT NOT NULL, report TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS address_candidates(
 id INTEGER PRIMARY KEY, import_hash TEXT NOT NULL REFERENCES address_imports(hash),
 source_row INTEGER NOT NULL, name TEXT NOT NULL, address TEXT NOT NULL,
 status TEXT NOT NULL, worker_id INTEGER REFERENCES workers(id),
 UNIQUE(import_hash,source_row));
'''
HEADERS = ('Naam', 'Straat', 'Nummer', 'Bus', 'Postcode', 'Gemeente')
FIELDS = ('street', 'number', 'unit', 'postal_code', 'city', 'country')


def cell_text(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        raise ValueError('Ongeldige adreswaarde.')
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def read_addresses(source):
    source = Path(source)
    content = source.read_bytes()
    if len(content) > 5_000_000:
        raise ValueError('Adresbestand groter dan 5 MB.')
    with ZipFile(source) as archive:
        if len(archive.infolist()) > 2000 or sum(e.file_size for e in archive.infolist()) > 64_000_000:
            raise ValueError('Adresbestand bevat te veel uitgepakte gegevens.')
    workbook = load_workbook(source, read_only=True, data_only=False)
    try:
        if 'Report' not in workbook.sheetnames:
            raise ValueError('Werkblad Report ontbreekt.')
        sheet = workbook['Report']
        rows = sheet.iter_rows()
        header = next(rows, ())
        if tuple(c.value for c in header[:6]) != HEADERS:
            raise ValueError('Verwachte adreskolommen ontbreken.')
        records = []
        for row_number, cells in enumerate(rows, 2):
            cells = cells[:6]
            if not any(c.value is not None for c in cells):
                continue
            if any(c.data_type in ('f', 'e') for c in cells):
                raise ValueError(f'Adresrij {row_number}: formule of Excel-fout niet toegestaan.')
            values = [cell_text(c.value) for c in cells]
            if len(values) != 6 or any(not values[i] for i in (0, 1, 2, 4, 5)):
                raise ValueError(f'Adresrij {row_number}: verplicht gegeven ontbreekt.')
            name = required(values[0])
            address = dict(zip(FIELDS[:5], values[1:]))
            address['country'] = ''  # Source contains no country; never infer it.
            for field in FIELDS[:5]:
                if field != 'unit':
                    required(address[field])
            records.append({'row': row_number, 'name': name, 'address': address})
        if not records:
            raise ValueError('Adresbestand bevat geen adressen.')
        return sha256(content).hexdigest(), records
    finally:
        workbook.close()


def append_address(db, worker_id, start, address, reason, source_row=None, source_hash=None):
    if type(worker_id) is not int or not db.execute('SELECT 1 FROM workers WHERE id=?', (worker_id,)).fetchone():
        raise ValueError('Onbekende werknemer.')
    start = valid_day(start)
    reason = required(reason)
    if not isinstance(address, dict):
        raise ValueError('Adres vereist.')
    cleaned = {}
    for field in FIELDS:
        value = address.get(field, '')
        if field in ('unit', 'country') and value == '':
            cleaned[field] = ''
        else:
            cleaned[field] = required(value)
    country = cleaned['country'].upper()
    if country and (len(country) != 2 or not country.isascii() or not country.isalpha()):
        raise ValueError('Land: gebruik een tweelettercode, bijvoorbeeld BE of NL.')
    cleaned['country'] = country
    last = db.execute('SELECT max(valid_from) FROM worker_addresses WHERE worker_id=?', (worker_id,)).fetchone()[0]
    if last and start < last:
        raise ValueError('Kies de laatste adresdatum of een latere datum. Historiek blijft bewaard.')
    return db.execute('INSERT INTO worker_addresses(worker_id,valid_from,address,changed_at,reason,source_row,source_hash) VALUES(?,?,?,?,?,?,?)',
        (worker_id, start, json.dumps(cleaned, ensure_ascii=False), datetime.now(timezone.utc).isoformat(), reason, source_row, source_hash)).lastrowid


def import_addresses(store, source, start, revision):
    start = valid_day(start)
    digest, records = read_addresses(source)
    counts = Counter(norm(r['name']) for r in records)
    with store.transaction(revision) as db:
        old = db.execute('SELECT * FROM address_imports WHERE hash=?', (digest,)).fetchone()
        if old:
            if old['valid_from'] != start:
                raise ValueError('Deze adreslijst is al ingelezen met een andere ingangsdatum.')
            return {**json.loads(old['report']), 'already_imported': True}
        workers = {r['name_key']: r['id'] for r in db.execute('SELECT * FROM workers')}
        report = {'rows': len(records), 'linked': 0, 'review': 0, 'valid_from': start}
        db.execute('INSERT INTO address_imports VALUES(?,?,?)', (digest, start, '{}'))
        for record in records:
            name_key = norm(record['name'])
            worker_id = workers.get(name_key)
            existing = worker_id and db.execute('SELECT 1 FROM worker_addresses WHERE worker_id=?', (worker_id,)).fetchone()
            status = 'LINKED' if worker_id and counts[name_key] == 1 and not existing else 'REVIEW'
            if status == 'LINKED':
                append_address(db, worker_id, start, record['address'], 'Adreslijst · ingangsdatum bevestigd door gebruiker', record['row'], digest)
                report['linked'] += 1
            else:
                report['review'] += 1
            db.execute('INSERT INTO address_candidates(import_hash,source_row,name,address,status,worker_id) VALUES(?,?,?,?,?,?)',
                (digest, record['row'], record['name'], json.dumps(record['address'], ensure_ascii=False), status, worker_id))
        db.execute('UPDATE address_imports SET report=? WHERE hash=?', (json.dumps(report), digest))
        return report


def snapshot(db):
    return {
        'addresses': [{**dict(r), 'address': json.loads(r['address']), 'geocode_status': 'NOT_REQUESTED'}
                      for r in db.execute('SELECT * FROM worker_addresses ORDER BY valid_from,id')],
        'address_candidates': [{**dict(r), 'address': json.loads(r['address'])}
                               for r in db.execute('SELECT c.*,i.valid_from FROM address_candidates c JOIN address_imports i ON i.hash=c.import_hash ORDER BY c.id')]
    }


def confirm_belgium(store, revision):
    """User-confirmed country only; never replace an explicit foreign country."""
    reason = 'Gebruiker bevestigt: alle werknemers wonen in België'
    report = {'profiles': 0, 'candidates': 0}
    with store.transaction(revision) as db:
        latest = {}
        for row in db.execute('SELECT * FROM worker_addresses ORDER BY valid_from,id'):
            latest[row['worker_id']] = dict(row)
        for row in latest.values():
            address = json.loads(row['address'])
            if address.get('country'):
                continue
            append_address(db, row['worker_id'], row['valid_from'],
                           {**address, 'country': 'BE'}, reason,
                           row['source_row'], row['source_hash'])
            report['profiles'] += 1
        for row in db.execute("SELECT * FROM address_candidates WHERE status='REVIEW'").fetchall():
            address = json.loads(row['address'])
            if address.get('country'):
                continue
            db.execute('UPDATE address_candidates SET address=? WHERE id=?',
                       (json.dumps({**address, 'country': 'BE'}, ensure_ascii=False), row['id']))
            report['candidates'] += 1
        db.execute('INSERT INTO matching_audit(changed_at,reason,details) VALUES(?,?,?)',
                   (datetime.now(timezone.utc).isoformat(), reason,
                    json.dumps({'action': 'address_country_confirmation', 'country': 'BE', **report})))
    return report


def apply(db, action, data):
    reason = required(data.get('reason'))
    if action == 'address_save':
        address = data.get('address')
        if not isinstance(address, dict) or not address.get('country'):
            raise ValueError('Bevestig eerst de landcode van het adres.')
        return append_address(db, data.get('worker_id'), data.get('valid_from'), address, reason)
    candidate = db.execute('SELECT c.*,i.valid_from FROM address_candidates c JOIN address_imports i ON i.hash=c.import_hash WHERE c.id=?', (data.get('candidate_id'),)).fetchone()
    if not candidate or candidate['status'] != 'REVIEW':
        raise ValueError('Adresbronregel is niet meer ter controle beschikbaar.')
    if action == 'address_link':
        worker_id = data.get('worker_id')
        append_address(db, worker_id, candidate['valid_from'], json.loads(candidate['address']), reason, candidate['source_row'], candidate['import_hash'])
        db.execute("UPDATE address_candidates SET status='LINKED',worker_id=? WHERE id=?", (worker_id, candidate['id']))
    elif action == 'address_ignore':
        db.execute("UPDATE address_candidates SET status='IGNORED' WHERE id=?", (candidate['id'],))
    else:
        raise ValueError('Onbekende adresactie.')
    db.execute('INSERT INTO matching_audit(changed_at,reason,details) VALUES(?,?,?)',
        (datetime.now(timezone.utc).isoformat(), reason, json.dumps({'action': action, 'candidate_id': candidate['id']})))
