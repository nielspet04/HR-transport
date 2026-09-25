"""Read-only Excel control export. Never request routes or alter HR configuration."""
from datetime import date
from io import BytesIO
import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from app.distances import whole_kms
from app.geocoding import fingerprint
from app.importers.reference import norm
from app.routing import plan, profile
from app.route_corrections import cached_distances,effective_distance


def text(value):
    # Export identifiers and labels as literal text, never executable formulas.
    value=str(value or '')
    return "'"+value if value.startswith(('=','+','-','@')) else value


def shift_label(shift):
    offset=shift.get('end_day_offset')
    if offset is None:offset=int(shift['end']<shift['start'])
    return f"{shift['start']} – {shift['end']}"+(f" (+{offset} dag)" if offset else '')


def export_rows(store, run_id):
    with store.connect() as db:
        config=store.matching_config(db)
        prepared=plan(db,config,run_id)
        workers={r['id']:r['name'] for r in config['workers']}
        cached={r['cache_key']:r for r in cached_distances(db)}
        homes={};sites={}
        for row in db.execute('SELECT * FROM worker_addresses'):
            homes.setdefault(fingerprint(json.loads(row['address'])),set()).add((row['worker_id'],row['valid_from']))
        for row in db.execute('SELECT * FROM location_address_versions'):
            sites.setdefault(fingerprint(json.loads(row['address'])),set()).add((row['location'],row['valid_from']))
        contexts={r['cache_key']:r['contexts'] for r in prepared['routes']}
        distance_rows=[]
        for key,r in cached.items():
            refs=contexts.get(key,[])
            if not refs:
                saved=db.execute('SELECT contexts FROM route_saved_contexts WHERE route_id=?',(r['id'],)).fetchone()
                if saved:refs=json.loads(saved['contexts'])
            if not refs:
                refs=[{'worker':workers.get(wid,'Onbekend'),'worker_id':wid,'location':location,'home_valid_from':hd,'location_valid_from':sd}
                      for wid,hd in sorted(homes.get(r['origin_hash'],set())) for location,sd in sorted(sites.get(r['destination_hash'],set()))]
            if not refs:refs=[{'worker':'Niet meer gekoppeld','location':'Niet meer gekoppeld','home_valid_from':'','location_valid_from':''}]
            seen=set()
            for c in refs:
                marker=(c['worker'],c.get('origin_location'),c['location'],c['home_valid_from'],c['location_valid_from'])
                if marker in seen:continue
                seen.add(marker)
                effective=effective_distance(r,c.get('worker_id'))
                message=('HR-correctie: '+effective['override_reason']+' · Mapbox: '+str(effective['mapbox_kms'])+' km') if effective['distance_source']=='HR' else r.get('message')
                distance_rows.append([text(c['worker']),'Fiets' if r['profile']=='cycling' else 'Auto',whole_kms(effective['kms']) if r['status']=='READY' else None,
                    text((c['origin_location']+' → ' if c.get('origin_location') else '')+c['location']),'HR-CORRECTIE' if effective['distance_source']=='HR' else r['status'],r['id'],date.fromisoformat(c['home_valid_from']) if c['home_valid_from'] else None,date.fromisoformat(c['location_valid_from']) if c['location_valid_from'] else None,text(message)])
        lookup={}
        for r in prepared['routes']:
            for c in r['contexts']:
                saved=cached.get(r['cache_key'])
                lookup[(c['worker_id'],norm(c['location']),c['day'],r['profile'])]=effective_distance(saved,c['worker_id']) if saved else None
        payloads={}
        for row in db.execute('SELECT payload FROM matching_runs ORDER BY id DESC'):
            p=json.loads(row['payload'])
            if p.get('source_sha256')==prepared['source_sha256'] and p['month'] in prepared['months']:payloads.setdefault(p['month'],p)
        shift_rows=[]
        for month,p in sorted(payloads.items()):
            agents={a['planet_id']:a for a in p['agents']}
            from app.shift_transport import choices as transport_choices
            overrides=transport_choices(db,p)
            for m in p['movements']:
                agent=agents[m['planet_id']];wid=agent.get('worker_id')
                name=workers.get(wid) or ' / '.join(agent.get('source_names',[])) or f"Planet {m['planet_id']}"
                choice=overrides.get(m.get('id'))
                selected_modes={'AUTO':'Privé auto','BIKE':'Fiets','TRAIN':'Trein','COMPANY_CAR':'Dienstwagen','MOBILITY_BUDGET':'Mob budget','TELEWORK':'Telework'}
                options=[{'mode':selected_modes[choice['mode']]}] if choice and choice['mode']!='DEFAULT' else m.get('routes') or [{'mode':'Niet ingesteld'}]
                for option in options:
                    mode=option['mode'];r=lookup.get((wid,norm(m.get('location') or ''),m['day'],profile(mode)))
                    excluded=norm(mode) in ('trein','dienstwagen','mob budget','telework')
                    km='n.v.t.' if excluded else whole_kms(r['kms']) if r and r['status']=='READY' else None
                    status='N.V.T.' if excluded else r['status'] if r else m['status'] if m['status']!='MATCHED' else 'AFSTAND ONTBREEKT'
                    if r and r.get('distance_source')=='HR':status='HR-CORRECTIE: '+r['override_reason']
                    sources=m.get('source_shifts',[])
                    shift=' / '.join(dict.fromkeys(shift_label(s) for s in sources))
                    customers=' / '.join(dict.fromkeys(s['customer'] for s in sources))
                    tasks=' / '.join(dict.fromkeys(s.get('task','') for s in sources))
                    shift_rows.append([text(name),text(mode),km,text(shift),date.fromisoformat(m['day']),text(customers),text(m.get('location') or m['source_location']),
                        text(m['planet_id']),', '.join(str(s['row']) for s in sources),status,r['id'] if r else None,'Vervoerskeuze controleren' if len(options)>1 else '',text(tasks)])
        return distance_rows,shift_rows,prepared['months']


def workbook_bytes(store, run_id):
    distances,shifts,months=export_rows(store,run_id)
    wb=Workbook();wb.remove(wb.active)
    definitions=[('Afstanden',['Naam','Vervoerstype','Enkele afstand (km)','Locatie','Status','Route-ID','Woonadres vanaf','Locatieadres vanaf','Melding'],distances,[30,20,22,30,25,12,20,20,65]),
                 ('Shiften',['Naam','Vervoerstype','Enkele afstand (km)','Shift','Datum','Customer','Locatie','Planet-ID','Bronregels','Status','Route-ID','Vervoerskeuze','Taak'],shifts,[30,20,22,42,15,42,30,15,20,30,12,32,35])]
    for name,headers,rows,widths in definitions:
        ws=wb.create_sheet(name);ws.sheet_view.showGridLines=False;ws.sheet_properties.tabColor='1F4E78'
        ws.cell(2,1,'Werknemersafstanden' if name=='Afstanden' else 'Shiftcontrole')
        ws.cell(2,1).font=Font(name='Arial',size=14,bold=True)
        ws.cell(3,1,'Enkele kilometers altijd naar boven afgerond. Geen uitbetalingsbestand.')
        ws.cell(4,1,'Alle opgeslagen Mapbox-afstanden, inclusief historische adresversies.' if name=='Afstanden' else 'Maanden: '+', '.join(months)+'. Eén rij per fase-2-beweging en vervoersoptie. Geen bevestigde vervoerskeuze.')
        ws.append(headers)  # row 5
        for values in rows:ws.append(values)
        for cells in ws.iter_rows(min_row=5):
            for cell in cells:
                cell.font=Font(name='Arial',size=10)
                cell.alignment=Alignment(vertical='center',horizontal='right' if isinstance(cell.value,(int,float)) else 'left')
                if cell.column==3:cell.number_format='0'
                if isinstance(cell.value,date):cell.number_format='dd/mm/yyyy'
                if name=='Shiften' and cell.row>5 and cell.column in (4,6,9,13):
                    cell.alignment=Alignment(vertical='center',wrap_text=True)
                    width=widths[cell.column-1]
                    ws.row_dimensions[cell.row].height=max(ws.row_dimensions[cell.row].height or 20,15*((len(str(cell.value or ''))//max(1,int(width)-3))+1))
        for cell in ws[5]:
            cell.fill=PatternFill('solid',fgColor='1F4E78');cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF');cell.alignment=Alignment(horizontal='center',vertical='center')
        for index,width in enumerate(widths,1):ws.column_dimensions[ws.cell(5,index).column_letter].width=width
        ws.freeze_panes='C6';ws.row_dimensions[5].height=25
        if rows:
            table=Table(displayName=name+'Controle',ref=f'A5:{ws.cell(5,len(headers)).column_letter}{ws.max_row}')
            table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
            ws.add_table(table)
        ws.auto_filter.ref=f'A5:{ws.cell(5,len(headers)).column_letter}{ws.max_row}'
    buffer=BytesIO();wb.save(buffer);return buffer.getvalue()
