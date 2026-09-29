"""Current employee address and route distances for technical reconciliation."""
from datetime import date
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.distances import whole_kms
from app.geocoding import fingerprint
from app.routing import profile


HEADERS = (
    'Werknemer naam', 'Adres', 'Mapbox km', 'Handmatig ingestelde km',
    'Locatie', 'Vervoerswijze',
)


def _text(value):
    value = str(value or '')
    return "'" + value if value.startswith(('=', '+', '-', '@')) else value


def _address(value):
    if not isinstance(value, dict):
        return ''
    street = ' '.join(part for part in (str(value.get('street') or '').strip(), str(value.get('number') or '').strip()) if part)
    unit = str(value.get('unit') or '').strip()
    locality = ' '.join(part for part in (str(value.get('postal_code') or '').strip(), str(value.get('city') or '').strip()) if part)
    country = str(value.get('country') or '').strip().upper()
    return ', '.join(part for part in (street + (f' bus {unit}' if unit else ''), locality, country) if part)


def _number(value):
    if value in (None, ''):
        return None
    try:
        result = Decimal(str(value).replace(',', '.'))
    except (InvalidOperation, ValueError, TypeError):
        return None
    if not result.is_finite():
        return None
    return int(result) if result == result.to_integral_value() else float(result)


def export_rows(store, as_of=None):
    day = as_of or date.today().isoformat()
    with store.connect() as db:
        workers = {row['id']: row['name'] for row in db.execute('SELECT id,name FROM workers')}
        homes = {}
        for row in db.execute('SELECT * FROM worker_addresses WHERE valid_from<=? ORDER BY valid_from,id', (day,)):
            homes[row['worker_id']] = dict(row)
        locations = {}
        for row in db.execute('SELECT * FROM location_address_versions WHERE valid_from<=? ORDER BY valid_from,id', (day,)):
            locations[row['location_key']] = dict(row)
        versions = {}
        for row in db.execute('SELECT * FROM versions WHERE valid_from<=? ORDER BY valid_from,id', (day,)):
            versions[row['route_id']] = dict(row)
        mapbox = {}
        for row in db.execute("SELECT * FROM route_distances WHERE status='READY' ORDER BY id"):
            mapbox[(row['origin_hash'], row['destination_hash'], row['profile'])] = dict(row)

        rows = []
        for route in db.execute('SELECT * FROM routes ORDER BY worker_id,location_key,mode_key,id'):
            home = homes.get(route['worker_id'])
            site = locations.get(route['location_key'])
            route_profile = profile(route['mode'])
            cached = None
            if home and site and route_profile:
                try:
                    cached = mapbox.get((fingerprint(json.loads(home['address'])),
                                         fingerprint(json.loads(site['address'])), route_profile))
                except (json.JSONDecodeError, TypeError, ValueError):
                    cached = None
            version = versions.get(route['id'])
            address = json.loads(home['address']) if home else None
            rows.append([
                _text(workers.get(route['worker_id'], 'Onbekende werknemer')),
                _text(_address(address)),
                whole_kms(cached['kms']) if cached and cached['kms'] is not None else None,
                _number(version['kms']) if version else None,
                _text(route['location']),
                _text(route['mode']),
            ])
        return rows


def workbook_bytes(store):
    rows = export_rows(store)
    workbook = Workbook(); sheet = workbook.active; sheet.title = 'Werknemerafstanden'
    sheet.sheet_view.showGridLines = False
    sheet['A1'] = 'Werknemers · Mapbox en handmatige kilometers'
    sheet['A1'].font = Font(name='Arial', size=14, bold=True)
    sheet['A2'] = 'Actuele woonadressen en route-instellingen op de exportdatum. Afstanden zijn enkele kilometers.'
    sheet.append(list(HEADERS))
    for row in rows:
        sheet.append(row)
    widths = (28, 48, 15, 24, 28, 20)
    for cell in sheet[3]:
        cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
        cell.alignment = Alignment(horizontal='center', vertical='center')
    for cells in sheet.iter_rows(min_row=4):
        for cell in cells:
            cell.font = Font(name='Arial', size=10)
            cell.alignment = Alignment(vertical='center', horizontal='right' if cell.column in (3, 4) else 'left')
            if cell.column in (3, 4): cell.number_format = '0.##'
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(3, index).column_letter].width = width
    sheet.freeze_panes = 'C4'; sheet.auto_filter.ref = f'A3:F{max(3, sheet.max_row)}'
    if rows:
        table = Table(displayName='Werknemerafstanden', ref=f'A3:F{sheet.max_row}')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        sheet.add_table(table)
    buffer = BytesIO(); workbook.save(buffer); return buffer.getvalue()
