"""Pl@net Total kms -> typed shifts, with explicit row accounting."""

from __future__ import annotations

import hashlib
import re
from collections import Counter
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.models.shift import ImportIssue, ImportReport, ImportResult, Shift


HEADERS = (
    "Id", "Last name", "First name", "Department", "Day", "Task",
    "Debut tache", "Fin tache", "Remark", "Customer", "Kms",
)
SHEET_NAME = "Total kms"
HEADER_SCAN_ROWS = 75


class ImportSourceError(ValueError):
    """Invalid source or structure; messages do not expose source values."""


class FieldError(ValueError):
    def __init__(self, code: str) -> None:
        super().__init__(code)
        self.code = code


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def header_key(value: Any) -> str:
    return re.sub(r"\s+", " ", value.strip()).casefold() if isinstance(value, str) else ""


def parse_text(value: Any, *, required: bool = True) -> str | None:
    if is_blank(value):
        if required:
            raise FieldError("MISSING_VALUE")
        return None
    if not isinstance(value, str):
        raise FieldError("INVALID_TEXT_TYPE")
    if value.startswith("="):
        raise FieldError("FORMULA_NOT_SUPPORTED")
    return value.strip()


def parse_day(value: Any) -> date:
    if is_blank(value):
        raise FieldError("MISSING_VALUE")
    if isinstance(value, datetime):
        if value.time() != time(0):
            raise FieldError("DATE_CONTAINS_TIME")
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value.strip()):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            pass
    raise FieldError("INVALID_DATE")


def parse_time(value: Any) -> time:
    if is_blank(value):
        raise FieldError("MISSING_VALUE")
    if isinstance(value, time) and value.tzinfo is None:
        return value
    if isinstance(value, str) and re.fullmatch(r"\d{2}:\d{2}(?::\d{2})?", value.strip()):
        try:
            return time.fromisoformat(value.strip())
        except ValueError:
            pass
    # Bare Excel serials/datetimes are not interpreted as times without a format.
    raise FieldError("INVALID_TIME")


def parse_kms(value: Any) -> Decimal | None:
    if is_blank(value):
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal, str)):
        raise FieldError("INVALID_KMS")
    if isinstance(value, str):
        value = value.strip()
        if value.startswith("="):
            raise FieldError("FORMULA_NOT_SUPPORTED")
        if not re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", value):
            raise FieldError("INVALID_KMS")
        value = value.replace(",", ".")
    try:
        result = Decimal(str(value))
    except InvalidOperation:
        raise FieldError("INVALID_KMS") from None
    if not result.is_finite():
        raise FieldError("INVALID_KMS")
    return result


def parse_end_time(value: Any) -> tuple[time, int | None]:
    # 24:00 explicitly encodes next midnight; other clock values have no offset.
    if isinstance(value, str) and value.strip() in {"24:00", "24:00:00"}:
        return time(0), 1
    return parse_time(value), None


def validate_month(month: str | None) -> None:
    if month is None:
        return
    if not isinstance(month, str) or not re.fullmatch(r"\d{4}-\d{2}", month):
        raise ImportSourceError("Maand moet het formaat YYYY-MM hebben.")
    try:
        date.fromisoformat(month + "-01")
    except ValueError:
        raise ImportSourceError("Ongeldige maand; gebruik YYYY-MM.") from None


def parse_shift(
    values: dict[str, Any], *, source_file: str, source_row: int
) -> tuple[Shift | None, list[ImportIssue]]:
    parsed: dict[str, Any] = {}
    issues: list[ImportIssue] = []
    text_fields = {
        "Id": "employee_id", "Last name": "last_name", "First name": "first_name",
        "Department": "department", "Task": "task", "Remark": "remark",
        "Customer": "customer",
    }
    for source_field, target_field in text_fields.items():
        try:
            parsed[target_field] = parse_text(
                values[source_field], required=source_field not in {"Department", "Remark"}
            )
        except FieldError as error:
            code = "INVALID_IDENTIFIER_TYPE" if source_field == "Id" and error.code == "INVALID_TEXT_TYPE" else error.code
            issues.append(ImportIssue(code, "ERROR", source_row, source_field))
    for source_field, target_field, parser in (
        ("Day", "day", parse_day),
        ("Debut tache", "start_time", parse_time),
        ("Kms", "kms", parse_kms),
    ):
        try:
            parsed[target_field] = parser(values[source_field])
        except FieldError as error:
            issues.append(ImportIssue(error.code, "ERROR", source_row, source_field))
    try:
        parsed["end_time"], parsed["end_time_day_offset"] = parse_end_time(values["Fin tache"])
    except FieldError as error:
        issues.append(ImportIssue(error.code, "ERROR", source_row, "Fin tache"))
    if issues:
        return None, issues
    shift = Shift(**parsed, source_file=source_file, source_sheet=SHEET_NAME, source_row=source_row)
    if shift.end_time_day_offset == 1:
        issues.append(ImportIssue("EXPLICIT_24H_END", "INFO", source_row, "Fin tache"))
    elif shift.end_time < shift.start_time:
        issues.append(ImportIssue("END_BEFORE_START", "WARNING", source_row, "Fin tache"))
    elif shift.end_time == shift.start_time:
        issues.append(ImportIssue("SAME_START_END_TIME", "WARNING", source_row, "Fin tache"))
    if shift.kms is None:
        issues.append(ImportIssue("MISSING_KMS", "WARNING", source_row, "Kms"))
    elif shift.kms < 0:
        issues.append(ImportIssue("NEGATIVE_KMS", "WARNING", source_row, "Kms"))
    return shift, issues


def import_planet(path: str | Path, *, month: str | None = None) -> ImportResult:
    """Validate the entire source, then optionally select one explicit month.

    Invalid rows remain visible in issues/report. No telework filtering, duplicate
    removal, employee/location matching, end-date inference or transport logic.
    """
    validate_month(month)
    path = Path(path)
    if not path.is_file() or path.suffix.lower() != ".xlsx":
        raise ImportSourceError("Bron moet een bestaand .xlsx-bestand zijn.")
    try:
        with path.open("rb") as stream:
            source_sha256 = hashlib.file_digest(stream, "sha256").hexdigest()
        workbook = load_workbook(path, read_only=True, data_only=False)
    except (OSError, BadZipFile, InvalidFileException, ValueError, KeyError, ParseError):
        raise ImportSourceError("Excelbron kan niet worden geopend.") from None
    try:
        if SHEET_NAME not in workbook.sheetnames:
            raise ImportSourceError("Verplicht werkblad 'Total kms' ontbreekt.")
        sheet = workbook[SHEET_NAME]
        # Derive actual bounds; dimension/filter metadata can omit valid rows.
        sheet.reset_dimensions()
        if next(sheet.iter_rows(), None) is None:
            raise ImportSourceError("Werkblad 'Total kms' is volledig leeg.")
        sheet.calculate_dimension(force=True)
        expected = {header_key(header): header for header in HEADERS}
        header_row = None
        mapping: dict[str, int] = {}
        header_values: tuple[Any, ...] = ()
        for row_number, row in enumerate(
            sheet.iter_rows(max_row=min(sheet.max_row, HEADER_SCAN_ROWS), values_only=True), 1
        ):
            keys = [header_key(value) for value in row]
            if not set(expected).issubset(keys):
                continue
            if any(count > 1 for key, count in Counter(keys).items() if key):
                raise ImportSourceError("Dubbele kolomnaam in de header.")
            header_row = row_number
            header_values = row
            mapping = {name: keys.index(key) for key, name in expected.items()}
            break
        if header_row is None:
            raise ImportSourceError("Geen volledige Pl@net-header gevonden in de eerste 75 rijen.")
        issues: list[ImportIssue] = []
        if any(header_key(value) and header_key(value) not in expected for value in header_values):
            issues.append(ImportIssue("EXTRA_COLUMNS_IGNORED", "WARNING", header_row))
        shifts: list[Shift] = []
        months: set[str] = set()
        seen = blanks = summaries = repeated = rejected = outside = 0
        unheaded = [index for index, value in enumerate(header_values) if is_blank(value)]
        for row_number, cells in enumerate(
            sheet.iter_rows(min_row=header_row + 1), header_row + 1
        ):
            row = tuple(cell.value for cell in cells)
            seen += 1
            if all(is_blank(value) for value in row):
                blanks += 1
                continue
            if all(header_key(row[index]) == header_key(name) for name, index in mapping.items()):
                repeated += 1
                issues.append(ImportIssue("REPEATED_HEADER", "WARNING", row_number))
                continue
            if any(not is_blank(row[index]) for index in unheaded):
                rejected += 1
                issues.append(ImportIssue("UNHEADED_DATA", "ERROR", row_number))
                continue
            values = {name: row[index] for name, index in mapping.items()}
            error_fields = [name for name, index in mapping.items() if cells[index].data_type == "e"]
            if error_fields:
                rejected += 1
                issues.extend(ImportIssue("SOURCE_CELL_ERROR", "ERROR", row_number, name) for name in error_fields)
                continue
            identifier = values["Id"]
            if isinstance(identifier, str) and re.fullmatch(r"(?:.+\s+)?Total", identifier.strip(), re.I):
                if all(is_blank(values[name]) for name in HEADERS if name not in {"Id", "Kms"}):
                    summaries += 1
                    continue
                rejected += 1
                issues.append(ImportIssue("AMBIGUOUS_SUMMARY_ROW", "ERROR", row_number, "Id"))
                continue
            shift, row_issues = parse_shift(values, source_file=path.name, source_row=row_number)
            issues.extend(row_issues)
            if shift is None:
                rejected += 1
                continue
            shift_month = shift.day.strftime("%Y-%m")
            months.add(shift_month)
            if month is not None and shift_month != month:
                outside += 1
            else:
                shifts.append(shift)
        if len(months) > 1:
            issues.append(ImportIssue("MULTIPLE_MONTHS", "WARNING"))
        if month is not None and month not in months:
            issues.append(ImportIssue("MONTH_NOT_PRESENT", "ERROR", field="month"))
        report = ImportReport(
            header_row, seen, blanks, summaries, repeated, rejected, len(shifts),
            outside, tuple(sorted(months)), month, source_sha256,
        )
        return ImportResult(tuple(shifts), tuple(issues), report)
    except (OSError, BadZipFile, ParseError):
        raise ImportSourceError("Excelinhoud kan niet volledig worden gelezen.") from None
    finally:
        workbook.close()
